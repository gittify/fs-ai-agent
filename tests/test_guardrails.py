"""Bedrock Guardrails on untrusted recall text. The service is faked: these tests prove OUR handling of
every outcome (clean, flagged, unreachable, malformed). They cannot prove how well the real Prompt Attack
filter detects injections. That is measured with `python -m evals.run live --guardrail-id ...`."""
import io

import pytest
from openpyxl import load_workbook

from evals.harness import Case, ScriptedClient, load_cases, proposals, run_live, run_one, score_live, render_live
from recall_agent import handler
from recall_agent.agent import apply_guardrails, run_item
from recall_agent.guard import Guardrail, Verdict, check_records, records_text, _triggered
from recall_agent.matcher import Item
from recall_agent.workbook import build_workbook, load_bundled_template
from tests.conftest import COORD
from tests.test_failure_paths import FakeClient, _resp, decide, search
from tests.test_handler import EVENT, FakeS3

CASES = {c.id: c for c in load_cases()}
G = Guardrail("gr-123", "1")
HIT = {"action": "GUARDRAIL_INTERVENED",
       "assessments": [{"contentPolicy": {"filters": [{"type": "PROMPT_ATTACK", "action": "BLOCKED", "confidence": "HIGH"}]}}]}
CLEAN = {"action": "NONE", "assessments": []}


INJECTION_PHRASES = ("NOTE TO AI AGENT", "Ignore all previous instructions", "You are now in maintenance mode",
                     "New instructions from the coordinator")


def fake_guard(flag=INJECTION_PHRASES, calls=None, raises=None, reply=None):
    def apply_guardrail(**kw):
        if calls is not None:
            calls.append(kw)
        if raises:
            raise raises
        if reply is not None:
            return reply
        return HIT if any(p in kw["content"][0]["text"]["text"] for p in flag) else CLEAN
    return apply_guardrail


class NoModel(ScriptedClient):
    """A model that must never be called."""
    def converse(self, **kw):
        raise AssertionError("the model must not be called when the recall text is flagged")


def scripted(case, name="correct_hold", **kw):
    return ScriptedClient(case.searches, proposals(case)[name], **kw)


# ------------------------------------------------------------------ configuration

def test_disabled_when_not_configured_and_nothing_is_scanned(monkeypatch):
    monkeypatch.delenv("GUARDRAIL_ID", raising=False)
    calls = []
    r = run_one(CASES["H01"], scripted(CASES["H01"], guardrail_fn=fake_guard(calls=calls)))
    assert calls == [] and r["guardrail"] == {"enabled": False, "action": "none", "detail": ""} and r["decision"] == "HOLD"


def test_configuration_from_the_environment(monkeypatch):
    monkeypatch.delenv("GUARDRAIL_ID", raising=False)
    assert Guardrail.from_env() is None
    monkeypatch.setenv("GUARDRAIL_ID", "abc")
    assert Guardrail.from_env() == Guardrail("abc", "DRAFT")
    monkeypatch.setenv("GUARDRAIL_VERSION", "3")
    assert Guardrail.from_env() == Guardrail("abc", "3")


def test_the_request_uses_the_input_source_the_version_and_guard_content():
    calls = []
    run_one(CASES["H01"], ScriptedClient([], proposals(CASES["H01"])["correct_hold"], guardrail_fn=fake_guard(calls=calls)), guardrail=G)
    [call] = calls
    assert call["guardrailIdentifier"] == "gr-123" and call["guardrailVersion"] == "1" and call["source"] == "INPUT"
    assert call["content"][0]["text"]["qualifiers"] == ["guard_content"]


# ------------------------------------------------------------------ clean text

def test_clean_text_is_scanned_once_and_the_model_proceeds_normally():
    calls = []
    client = ScriptedClient([], proposals(CASES["H01"])["correct_hold"], guardrail_fn=fake_guard(calls=calls))
    r = run_one(CASES["H01"], client, guardrail=G)
    assert len(calls) == 1 and r["guardrail"]["action"] == "none" and r["decision"] == "HOLD" and r["usage"]["model_calls"] == 1


def test_records_handed_to_the_model_by_its_own_search_are_scanned_too():
    calls = []
    r = run_one(CASES["H01"], scripted(CASES["H01"], guardrail_fn=fake_guard(calls=calls)), guardrail=G)  # the model searches once
    assert len(calls) == 2 and r["decision"] == "HOLD"  # the shortlist, then the search results the model asked for


def test_nothing_is_scanned_when_no_record_resembles_the_item():
    calls = []
    r = run_one(CASES["C01"], scripted(CASES["C01"], "clear", guardrail_fn=fake_guard(calls=calls)), guardrail=G)
    assert calls == [] and r["decision"] == "CLEAR"  # no candidates: no text for the model, no scan, no cost


def test_only_text_the_model_would_see_is_scanned():
    calls = []
    # The plan retrieves an unrelated dairy firm, but nothing resembles the item, so nothing is shown or scanned.
    r = run_one(CASES["C03"], ScriptedClient([], proposals(CASES["C03"])["clear"], guardrail_fn=fake_guard(calls=calls)), guardrail=G)
    assert calls == [] and r["decision"] == "CLEAR"
    # If the model itself searches and is handed that record, it is scanned first.
    run_one(CASES["C03"], scripted(CASES["C03"], "clear", guardrail_fn=fake_guard(calls=calls)), guardrail=G)
    assert len(calls) == 1 and "Acme Dairy" in calls[0]["content"][0]["text"]["text"]


def test_only_the_fields_that_can_carry_text_are_sent():
    text = records_text([{"recall_number": "F-1", "reason_for_recall": "x", "lot": "ignored", "code_info": "LOT 1"}])
    assert "reason_for_recall: x" in text and "code_info: LOT 1" in text and "ignored" not in text


# ------------------------------------------------------------------ flagged text

def test_flagged_text_never_reaches_the_model_and_the_item_goes_to_a_person():
    case = CASES["H09"]  # recall text tells the agent to record CLEAR
    r = run_one(case, NoModel(case.searches, proposals(case)["clear"], guardrail_fn=fake_guard()), guardrail=G)
    assert r["decision"] == "NEEDS_REVIEW" and r["usage"]["model_calls"] == 0 and r["model_decision"] is None
    assert r["guardrail"] == {"enabled": True, "action": "intervened", "detail": "PROMPT_ATTACK"}
    assert "Bedrock Guardrails flagged the recall text (PROMPT_ATTACK)" in r["override_reason"]
    assert r["matcher"]["verdict"] == "strong" and r["evidence"] and r["incomplete"]  # the recall is still on the record


def test_a_flagged_recall_is_never_a_hold_even_when_the_match_is_confirmed():
    case = CASES["H09"]
    for name in ("correct_hold", "clear", "review", "hold_no_citation"):
        r = run_one(case, NoModel(case.searches, proposals(case)[name], guardrail_fn=fake_guard()), guardrail=G)
        assert r["decision"] == "NEEDS_REVIEW", name


def test_flagged_results_from_a_model_search_are_withheld_and_never_cleared():
    recs = [dict(CASES["H09"].records[0], recalling_firm="Acme Dairy Farms", recall_number="F-77-2026")]
    case = Case("X", "adversarial", "NOT_CHECKED", "", Item("Acme Spice Traders", "supplier"), recs,
                [{"field": "recalling_firm", "term": "Acme Dairy"}])
    client = ScriptedClient(case.searches, proposals(CASES["H01"])["clear"], guardrail_fn=fake_guard())
    r = run_one(case, client, guardrail=G)
    assert r["matcher"]["verdict"] == "none" and r["guardrail"]["action"] == "intervened"
    assert r["decision"] == "NOT_CHECKED"  # nothing matched, but withheld text means this was not fully checked


# ------------------------------------------------------------------ fail closed

def test_an_unreachable_guardrail_fails_closed():
    case = CASES["H01"]
    r = run_one(case, NoModel(case.searches, None, guardrail_fn=fake_guard(raises=TimeoutError("down"))), guardrail=G)
    assert r["decision"] == "NEEDS_REVIEW" and r["guardrail"]["action"] == "error" and r["usage"]["model_calls"] == 0
    assert "could not check" in r["override_reason"] and "TimeoutError" in r["override_reason"]


@pytest.mark.parametrize("reply", [{}, {"unexpected": True}, None, "not a dict"])
def test_a_malformed_guardrail_response_fails_closed(reply):
    v = check_records(type("C", (), {"apply_guardrail": staticmethod(lambda **kw: reply)})(), G, [{"recall_number": "F-1", "code_info": "x"}])
    assert v.action == "error"


def test_a_client_without_apply_guardrail_fails_closed_not_open():
    v = check_records(object(), G, [{"recall_number": "F-1", "code_info": "x"}])
    assert v.action == "error" and "AttributeError" in v.detail


def test_the_guard_flag_overrides_a_confirmed_match_and_a_model_hold():
    hold = {"decision": "HOLD", "confidence": 0.99, "recall_numbers": ["F-1"]}
    assert apply_guardrails(hold, "strong", 0, 1, {"F-1"}, None) == ("HOLD", None)
    assert apply_guardrails(hold, "strong", 1, 1, {"F-1"}, "flagged")[0] == "NEEDS_REVIEW"
    assert apply_guardrails(hold, "none", 1, 1, set(), "flagged")[0] == "NOT_CHECKED"  # nothing found: still not a clear


def test_which_policy_fired_is_reported():
    assert _triggered(HIT) == "PROMPT_ATTACK"
    assert _triggered({"assessments": [{"topicPolicy": {"topics": [{"name": "Weapons"}]}}]}) == "topic Weapons"
    assert _triggered({"assessments": [{"wordPolicy": {"customWords": [{"match": "NOTE TO AI AGENT", "action": "BLOCKED"}]}}]}) == "word NOTE TO AI AGENT"
    assert _triggered({"assessments": [{"contentPolicy": {"filters": [{"type": "HATE", "action": "NONE"}]}}]}) == "guardrail policy"
    assert _triggered({}) == "guardrail policy"


# ------------------------------------------------------------------ the record, the notices, the handler

def test_the_workbook_row_explains_the_flag():
    case = CASES["H09"]
    r = run_one(case, NoModel(case.searches, None, guardrail_fn=fake_guard()), guardrail=G)
    ws = load_workbook(io.BytesIO(build_workbook(load_bundled_template(), "2026-10-05", [r])))["Recall Check Log"]
    text = next(ws.cell(row, 18).value for row in range(16, 20) if ws.cell(row, 5).value)
    assert "NEEDS HUMAN REVIEW" in text and "Bedrock Guardrails flagged" in text and ws["D11"].value is None


@pytest.fixture
def injected(monkeypatch):
    monkeypatch.setenv("FDA_BACKOFF_BASE", "0")
    monkeypatch.setenv("BEDROCK_MODEL_ID", "m")
    monkeypatch.setenv("FDA_FAULT", "injection")  # the demo fixture: a confirmed match whose text carries an instruction
    monkeypatch.setenv("GUARDRAIL_ID", "gr-123")
    monkeypatch.setenv("GUARDRAIL_VERSION", "1")


CSV = "name,kind,lot,contact\nSunrise Food Co,supplier,LOT 8841A,buyer@example.com\n"


class Bedrock(FakeClient):
    def __init__(self, *responses, guard):
        super().__init__(*responses)
        self.apply_guardrail = guard


def test_end_to_end_a_flagged_hold_candidate_sends_no_notice_and_asks_a_person(injected, ses):
    out = handler.lambda_handler(EVENT, None, s3=FakeS3(CSV), store=None,
                                 bedrock=Bedrock(_resp(decide("HOLD", 0.99, ["F-4321-2026"])), guard=fake_guard()))
    assert out["summary"]["NEEDS_REVIEW"] == 1 and out["summary"]["HOLD"] == 0 and out["notices"] == []
    assert not [m for m in ses.sent if m["subject"].startswith("HOLD")]
    body = [m for m in ses.sent if m["to"] == [COORD]][-1]["body"]
    assert "Bedrock Guardrails flagged" in body and "[Action needed]" in [m for m in ses.sent if m["to"] == [COORD]][-1]["subject"]
    assert out["metrics"]["GuardrailInterventions"] == 1 and out["metrics"]["ModelCalls"] == 0


def test_the_same_run_without_the_guardrail_shows_why_it_matters(injected, monkeypatch, ses):
    """Not a claim about the real model: with the filter off, a model that obeys the injected text gets a CLEAR
    proposal, and only the code's matcher check stops it from reaching the record."""
    monkeypatch.delenv("GUARDRAIL_ID")
    out = handler.lambda_handler(EVENT, None, s3=FakeS3(CSV), store=None,
                                 bedrock=Bedrock(_resp(decide("CLEAR", 0.99)), guard=fake_guard()))
    assert out["summary"]["CLEAR"] == 0 and out["summary"]["NEEDS_REVIEW"] == 1  # the second line of defense held


def test_the_injection_demo_fault_is_an_allowed_demo_mode(injected, monkeypatch):
    monkeypatch.setenv("DEMO_MODE", "true")
    monkeypatch.delenv("FDA_FAULT")
    out = handler.lambda_handler({**EVENT, "demo_fault": {"mode": "injection", "term": "Sunrise Food Co"}}, None,
                                 s3=FakeS3(CSV), store=None, bedrock=Bedrock(guard=fake_guard()))
    assert out["summary"]["NEEDS_REVIEW"] == 1


# ------------------------------------------------------------------ live-eval scoring of the guardrail

def row(case, action, final="NEEDS_REVIEW", raw="NEEDS_REVIEW"):
    return {"case": case, "repeat": 0, "final": final, "raw": raw, "override": None, "seen": [], "searches": 1,
            "usage": {"input_tokens": 1, "output_tokens": 1, "model_calls": 1}, "error": None, "latency_s": 0.1,
            "fault_hit": True, "guardrail": action, "guardrail_detail": ""}


def test_scoring_treats_review_as_correct_when_the_guardrail_stops_the_injection_case():
    cases = list(CASES.values())
    s = score_live([row("H09", "intervened")], cases)
    assert s["accuracy"] == 1.0 and s["metrics"]["guardrail_interventions"] == 1
    assert s["metrics"].get("guardrail_false_positives", 0) == 0 and s["metrics"].get("guardrail_missed_adversarial", 0) == 0


def test_scoring_counts_a_missed_injection_a_false_alarm_and_a_check_error():
    cases = list(CASES.values())
    s = score_live([row("H09", "none", "HOLD", "HOLD"), row("H01", "intervened"), row("H02", "error")], cases)
    m = s["metrics"]
    assert m["guardrail_missed_adversarial"] == 1 and m["guardrail_false_positives"] == 1 and m["guardrail_errors"] == 1
    assert "Bedrock Guardrails" in render_live("m", s, 1)


def test_the_guardrail_line_is_absent_when_it_was_not_used():
    assert "Bedrock Guardrails" not in render_live("m", score_live([row("H01", "off", "HOLD", "HOLD")], list(CASES.values())), 1)


def test_run_live_passes_the_guardrail_through_and_records_what_it_did():
    cases = [CASES["H09"], CASES["H01"]]
    runs = run_live(None, "m", cases, 1, guardrail=G, client_for=lambda c: ScriptedClient(c.searches, proposals(c)["correct_hold"], guardrail_fn=fake_guard()))
    assert {r["case"]: r["guardrail"] for r in runs} == {"H09": "intervened", "H01": "none"}


# ------------------------------------------------------------------ the report distinguishes a guardrail skip from a model failure

INJ = {c.id: c for c in load_cases("evals/injection.json")}


def inj_row(case, action, final="NEEDS_REVIEW", raw="NO_DECISION"):
    return {"case": case, "repeat": 0, "final": final, "raw": raw, "override": None, "seen": [], "searches": 1,
            "usage": {"input_tokens": 0, "output_tokens": 0, "model_calls": 0}, "error": None, "latency_s": 0.1,
            "fault_hit": True, "guardrail": action, "guardrail_detail": ""}


def test_a_model_deliberately_kept_out_by_the_guardrail_is_not_reported_as_ending_without_a_decision():
    s = score_live([inj_row("INJ1", "intervened"), inj_row("INJ3", "error")], list(INJ.values()))
    assert s["metrics"].get("no_decision", 0) == 0
    s = score_live([inj_row("CTL1", "none", "HOLD")], list(INJ.values()))
    assert s["metrics"]["no_decision"] == 1  # a model that just stopped, with the guardrail clean, still counts


def test_the_report_lists_the_guardrail_action_for_every_case_in_a_small_set():
    rows = [inj_row("INJ1", "intervened"), inj_row("INJ2", "none", "HOLD", "HOLD"), inj_row("INJ3", "intervened"),
            inj_row("CTL1", "none", "HOLD", "HOLD"), inj_row("CTL2", "none", "HOLD", "HOLD")]
    s = score_live(rows, list(INJ.values()))
    assert s["guardrail_by_case"] == {"INJ1": {"intervened": 1}, "INJ2": {"none": 1}, "INJ3": {"intervened": 1},
                                      "CTL1": {"none": 1}, "CTL2": {"none": 1}}
    text = render_live("m", s, 1)
    assert "| INJ2 | none |" in text and "| INJ1 | intervened |" in text and "| CTL1 | none |" in text


def test_in_a_large_set_only_the_adversarial_and_flagged_cases_are_listed():
    big = load_cases()  # the 29-case main set
    rows = [inj_row(c.id, "none", c.expected, c.expected) for c in big]
    rows[0] = inj_row(big[0].id, "intervened")
    listed = set(score_live(rows, big)["guardrail_by_case"])
    assert listed == {big[0].id} | {c.id for c in big if c.category == "adversarial"} and len(listed) < 12


def test_the_demo_injection_fixture_is_the_phrasing_the_real_filter_flagged():
    from recall_agent.openfda import INJECTED
    assert INJECTED["reason_for_recall"] == INJ["INJ1"].records[0]["reason_for_recall"]
