"""Workbook writer tests, run against the coordinator's real template."""
import io

import pytest
from openpyxl import load_workbook

from recall_agent import handler, openfda, workbook
from recall_agent.matcher import Item
from recall_agent.workbook import AGENT_NAME, AUDIT_SHEET, WorkbookError, build_workbook, load_bundled_template
from tests.conftest import COORD
from tests.test_failure_paths import FakeClient, _resp, decide, search
from tests.test_handler import EVENT, FakeS3

RUN = "2026-10-05"
REC = {**openfda.LOOKALIKE}
CALL = {"field": "recalling_firm", "term": "Sunrise Food Co", "status": "matches", "records": 1,
        "attempts": 1, "error": None, "url": "https://api.fda.gov/food/enforcement.json?search=x"}


def hold(name="Sunrise Food Co"):
    return {"item": name, "kind": "supplier", "decision": "HOLD", "override_reason": None,
            "model_decision": {"decision": "HOLD", "confidence": 0.92, "recall_numbers": ["F-1234-2026"],
                               "reasoning": "Firm and lot match."},
            "matcher": {"verdict": "strong", "matches": [
                {"recall_number": "F-1234-2026", "classification": "Class I", "identity": "exact",
                 "lot_overlap": ["LOT 8841A"], "level": "strong"}]},
            "evidence": [REC], "tool_calls": [CALL], "model_error": None}


def review(name="Sunrise Foods Inc"):
    r = hold(name)
    r.update(decision="NEEDS_REVIEW", override_reason="model said HOLD but matcher found no confirmed lot match")
    r["matcher"] = {"verdict": "possible", "matches": [
        {"recall_number": "F-1234-2026", "classification": "Class I", "identity": "close",
         "lot_overlap": [], "level": "possible"}]}
    return r


def clear(name="Acme Spice Traders"):
    return {"item": name, "kind": "supplier", "decision": "CLEAR", "override_reason": None,
            "model_decision": {"decision": "CLEAR", "confidence": 0.95, "reasoning": "No match."},
            "matcher": {"verdict": "none", "matches": []}, "evidence": [],
            "tool_calls": [{**CALL, "term": name, "status": "clear", "records": 0}], "model_error": None}


def unchecked(name="Globex Imports", reason="a search or model call failed; item could not be verified"):
    r = handler.not_checked(Item(name, "supplier"), reason)
    r["tool_calls"] = [{**CALL, "term": name, "status": "error", "records": 0, "attempts": 4,
                        "error": "gave up after 4 attempts: HTTP 429"}]
    return r


SENT = {"item": "Sunrise Food Co", "status": "sent", "to": "purchasing@example.com", "recalls": ["F-1234-2026"], "error": ""}


def build(results, notices=(), notes=()):
    data = build_workbook(load_bundled_template(), RUN, results, notices, notes)
    return load_workbook(io.BytesIO(data))


def cell(wb, label):
    ws = wb["Recall Check Log"]
    r = next(r for r in range(1, 15) if str(ws.cell(r, 3).value).strip() == label)
    return ws.cell(r, 4).value


def table(wb):
    ws = wb["Recall Check Log"]
    heads = {c.value: c.column for c in ws[15] if c.value}
    out = []
    for r in range(16, 48):
        if any(ws.cell(r, c).value for c in heads.values()):
            out.append({h: ws.cell(r, c).value for h, c in heads.items()})
    return out


def test_header_block_and_human_signoff_left_blank():
    wb = build([hold()], [SENT])
    assert cell(wb, "Check Date:") == RUN and cell(wb, "Checked By:") == AGENT_NAME
    assert cell(wb, "Next Review Date:") == "2026-10-12"
    assert cell(wb, "Signoff (Name & Date):") is None
    assert cell(wb, "No Applicable Recalls Found:") is None
    assert "not checked by the agent" in wb["Recall Check Log"]["D7"].value


def test_delivered_hold_row_is_complete_and_attributed():
    [row] = table(build([hold()], [SENT]))
    assert row["Recall Notice ID / Event ID"] == "F-1234-2026 / 98765"
    assert row["Recall Classification"] == "Class I" and row["Action Taken"] == "Hold"
    assert row["Reviewer (Name & Date)"] == f"{AGENT_NAME}, {RUN}" and row["Status"] == "Open"
    assert row["Lot / Batch / Date Codes"] == "LOT 8841A" and row["Source / URL"].startswith("https://api.fda.gov")
    assert "last 365 days" in row["Source Checked & Filter Used"]
    assert "Hold notice sent to purchasing@example.com" in row["Applicability Determination"]
    # cells with no FDA-side data are left for the inventory/human
    for h in ("UPC / Product Code", "Affected SKUs", "Qty Received", "Qty On Hand", "Customer Impact", "Disposition"):
        assert row[h] is None


def test_undelivered_hold_does_not_claim_a_hold_happened():
    failed = {**SENT, "status": "failed", "error": "MessageRejected"}
    [row] = table(build([hold()], [failed]))
    assert row["Action Taken"] is None and row["Reviewer (Name & Date)"] is None
    assert "NOT DELIVERED" in row["Applicability Determination"]


def test_review_and_not_checked_rows_leave_reviewer_blank_and_block_the_all_clear_box():
    wb = build([review(), unchecked(), clear()])
    rows = table(wb)
    assert [r["Recall Notice ID / Event ID"] for r in rows] == ["F-1234-2026 / 98765", "NOT CHECKED"]
    assert all(r["Reviewer (Name & Date)"] is None and r["Action Taken"] is None for r in rows)
    assert "NEEDS HUMAN REVIEW" in rows[0]["Applicability Determination"]
    assert "NOT CHECKED, could not verify" in rows[1]["Applicability Determination"]
    assert rows[1]["Supplier / Manufacturer"] == "Globex Imports"
    assert cell(wb, "No Applicable Recalls Found:") is None


def test_all_clear_week_ticks_the_box_adds_no_rows_and_lists_everything_on_the_audit_sheet():
    wb = build([clear("Acme Spice Traders"), clear("Initech Foods")])
    assert table(wb) == [] and cell(wb, "No Applicable Recalls Found:") == f"Yes ({AGENT_NAME})"
    audit = wb[AUDIT_SHEET]
    assert [audit.cell(r, 1).value for r in (5, 6)] == ["Acme Spice Traders", "Initech Foods"]
    assert "clear" in audit.cell(5, 5).value


def test_evidence_link_points_at_the_matching_audit_row():
    results = [clear(), hold()]  # the hold is the second item, so audit row 6
    wb = build(results, [SENT])
    ws = wb["Recall Check Log"]
    link = ws.cell(16, 19).hyperlink
    assert link.location == f"'{AUDIT_SHEET}'!A6" and wb[AUDIT_SHEET].cell(6, 1).value == "Sunrise Food Co"


def test_every_dropdown_value_written_is_in_the_templates_own_list():
    wb = build([hold(), review(), unchecked(), clear()], [SENT])
    ws = wb["Recall Check Log"]
    allowed = {}
    for dv in ws.data_validations.dataValidation:
        col = next(iter(dv.sqref.ranges)).min_col
        allowed[ws.cell(15, col).value] = dv.formula1.strip('"').split(",")
    assert set(allowed) == {"Recall Classification", "Action Taken", "Status"}
    for row in table(wb):
        for header, vals in allowed.items():
            assert row[header] is None or row[header] in vals


def test_overflow_fills_the_table_ends_with_a_plus_n_row_and_never_ticks_the_box():
    wb = build([unchecked(f"Supplier {i}") for i in range(40)])
    rows = table(wb)
    assert len(rows) == 32 and rows[-1]["Recall Notice ID / Event ID"] == "+9 MORE"
    assert wb["Recall Check Log"]["B48"].value == "Notes:"  # the Notes block was not overwritten
    assert cell(wb, "No Applicable Recalls Found:") is None
    assert wb[AUDIT_SHEET].max_row >= 5 + 40 - 1


def test_dropdown_ranges_no_longer_run_into_the_notes_block():
    ws = build([clear()])["Recall Check Log"]
    assert {str(dv.sqref) for dv in ws.data_validations.dataValidation} == {"K16:K47", "U16:U47", "P16:P47"}


def test_template_formatting_and_other_sheets_are_untouched():
    original = load_workbook(io.BytesIO(load_bundled_template()))
    wb = build([hold()], [SENT])
    assert wb.sheetnames == original.sheetnames + [AUDIT_SHEET]
    a, b = original["Recall Check Log"]["B15"], wb["Recall Check Log"]["B15"]
    assert a.fill.fgColor.rgb == b.fill.fgColor.rgb and a.font.b == b.font.b
    for name in ("FSVP Verification", "Instructions"):
        old = [[c.value for c in row] for row in original[name].iter_rows()]
        new = [[c.value for c in row] for row in wb[name].iter_rows()]
        assert old == new


def test_edited_template_fails_loudly_instead_of_guessing():
    wb = load_workbook(io.BytesIO(load_bundled_template()))
    wb["Recall Check Log"]["U15"] = "Case Status"  # coordinator renamed a column
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(WorkbookError) as e:
        build_workbook(buf.getvalue(), RUN, [hold()], [SENT])
    assert "Status" in str(e.value)


def test_dropdown_value_missing_from_template_list_is_an_error():
    wb = load_workbook(io.BytesIO(load_bundled_template()))
    ws = wb["Recall Check Log"]
    for dv in ws.data_validations.dataValidation:
        if dv.formula1 and "Quarantine" in dv.formula1:
            dv.formula1 = '"Quarantine,Return,Destroy"'  # "Hold" removed
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(WorkbookError):
        build_workbook(buf.getvalue(), RUN, [hold()], [SENT])


# ---- handler integration -------------------------------------------------------------------

CSV = "name,kind,lots,aliases,contact\nSunrise Food Co,supplier,LOT 8841A,,purchasing@example.com\n"


@pytest.fixture(autouse=True)
def env(monkeypatch):
    monkeypatch.setenv("FDA_BACKOFF_BASE", "0")
    monkeypatch.setenv("BEDROCK_MODEL_ID", "test-model")
    monkeypatch.setenv("FDA_FAULT", "lookalike")


def run_handler(monkeypatch=None):
    s3 = FakeS3(CSV)
    turns = [_resp(search("Sunrise Food Co")), _resp(decide("HOLD", 0.92, ["F-1234-2026"]))]
    out = handler.lambda_handler(EVENT, None, s3=s3, bedrock=FakeClient(*turns), store=None)
    return out, s3


def test_handler_saves_the_workbook_and_the_summary_points_at_it(ses):
    out, s3 = run_handler()
    key = ("bkt", "records/2026-10-05/weekly-recall-check.xlsx")
    assert out["workbook_key"] == key[1] and key in s3.objects
    wb = load_workbook(io.BytesIO(s3.objects[key]))
    assert table(wb)[0]["Action Taken"] == "Hold"
    assert "weekly-recall-check.xlsx" in [m for m in ses.sent if m["to"] == [COORD]][-1]["body"]


def test_broken_template_degrades_to_markdown_and_json_and_says_so(monkeypatch, ses):
    monkeypatch.setattr(handler.workbook, "load_bundled_template", lambda: b"not a workbook")
    out, s3 = run_handler()
    assert out["workbook_key"] is None
    assert ("bkt", "records/2026-10-05/weekly-recall-check.md") in s3.objects
    assert not any(k[1].endswith(".xlsx") for k in s3.objects)
    assert "Workbook could not be generated" in [m for m in ses.sent if m["to"] == [COORD]][-1]["body"]
