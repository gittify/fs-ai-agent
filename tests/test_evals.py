"""Tests for the eval harness itself, and regression tests for what the evals found."""
import io
import json
import os
import re
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from evals import run as eval_cli
from evals.harness import (Case, ScriptedClient, Transport, emulated_openfda, load_cases, proposals,
                           render_live, render_offline, run_live, run_offline, run_one, score_live, wilson,
                           _phrase_in)
from recall_agent import agent, openfda
from recall_agent.agent import apply_guardrails
from recall_agent.workbook import build_workbook, load_bundled_template
from recall_agent.notify import summary_email

CASES = load_cases()
BY_ID = {c.id: c for c in CASES}


# ------------------------------------------------------------------ the case set itself

def test_case_set_is_broad_unique_and_every_decision_is_represented():
    assert len(CASES) >= 25 and len({c.id for c in CASES}) == len(CASES)
    assert {c.expected for c in CASES} == {"HOLD", "CLEAR", "NEEDS_REVIEW", "NOT_CHECKED"}
    assert {c.category for c in CASES} >= {"hold", "review", "clear", "tool_failure", "adversarial"}


def test_hold_labels_are_independently_consistent_with_the_policy():
    """Guard against a mislabeled fixture: a HOLD case must really contain one of our lots in a cited record."""
    norm = lambda s: re.sub(r"[^a-z0-9]", "", s.lower())
    for c in CASES:
        if c.expected != "HOLD":
            continue
        cited = [r for r in c.records if r["recall_number"] in c.cite]
        ours = [norm(l.code) for l in c.item.lot_info]
        assert any(l in norm(r["code_info"]) for r in cited for l in ours), c.id


def test_invalid_case_files_are_rejected_with_all_problems_listed(tmp_path):
    bad = {"cases": [
        {"id": "X1", "category": "hold", "expected": "HOLD", "description": "", "searches": [], "cite": [],
         "item": {"name": "A", "kind": "supplier", "lots": []}, "fda": {"records": []}},
        {"id": "X1", "category": "clear", "expected": "MAYBE", "description": "", "searches": [{}], "cite": [],
         "item": {"name": "A", "kind": "supplier", "lots": []}, "fda": {"records": [], "fault": "meteor"}}]}
    path = tmp_path / "bad.json"
    path.write_text(json.dumps(bad))
    with pytest.raises(ValueError) as e:
        load_cases(path)
    msg = str(e.value)
    assert "duplicate id" in msg and "expected must be one of" in msg and "unknown fault" in msg and "must cite" in msg


# ------------------------------------------------------------------ the emulator

def test_phrase_matching_is_contiguous_and_case_and_punctuation_insensitive():
    assert _phrase_in("sunrise foods", "SUNRISE FOODS, LLC.")
    assert not _phrase_in("sunrise inc", "Sunrise Foods Inc")  # words must be adjacent, like a phrase search
    assert not _phrase_in("", "anything")


def test_emulator_answers_from_the_corpus_and_honors_fault_on():
    case = BY_ID["T07"]  # fault only on product_description searches
    t = Transport(case)
    ok = t('https://x/?search=recalling_firm:"Sunrise+Foods+Inc"+AND+report_date:[1+TO+2]&limit=10')
    assert ok[0] == 200 and json.loads(ok[1])["results"][0]["recall_number"] == "F-27-2026"
    with pytest.raises(TimeoutError):
        t('https://x/?search=product_description:"sesame+tahini"+AND+report_date:[1+TO+2]&limit=10')
    miss = Transport(BY_ID["C01"])('https://x/?search=recalling_firm:"Nobody"+AND+report_date:[1+TO+2]')
    assert miss[0] == 404 and "NOT_FOUND" in miss[1]


def test_emulation_restores_the_real_transport_and_environment(monkeypatch):
    monkeypatch.setenv("FDA_FAULT", "not_found")
    monkeypatch.delenv("FDA_BACKOFF_BASE", raising=False)
    real = openfda._request
    with emulated_openfda(BY_ID["C01"]):
        assert openfda._request is not real and "FDA_FAULT" not in os.environ
    assert openfda._request is real and os.environ["FDA_FAULT"] == "not_found" and "FDA_BACKOFF_BASE" not in os.environ


# ------------------------------------------------------------------ the offline eval

def test_offline_eval_has_zero_violations():
    report = run_offline(CASES)
    assert report["violations"] == [], report["violations"][:5]
    assert report["runs"] == len(CASES) * report["proposals"]
    assert all(t["applicable"] > 0 for t in report["invariants"].values())  # no invariant is vacuous


def test_the_eval_can_fail_a_naive_guardrail_lets_the_model_decide(monkeypatch):
    """Mutation test: with the guardrails replaced by 'trust the model', the eval must light up."""
    def naive(model_decision, verdict, errors, ok_searches, seen, guard_flag=None):
        d = (model_decision or {}).get("decision") if isinstance(model_decision, dict) else None
        return (d, None) if d in ("HOLD", "CLEAR", "NEEDS_REVIEW") else ("NEEDS_REVIEW", "no decision")
    monkeypatch.setattr(agent, "apply_guardrails", naive)
    inv = {t: v["violations"] for t, v in run_offline(CASES)["invariants"].items()}
    assert inv["SAFE-1"] > 0 and inv["SAFE-2"] > 0 and inv["SAFE-3"] > 0


def test_the_eval_catches_the_bug_it_found_errors_hiding_a_found_recall(monkeypatch):
    """Regression: the original rule turned ANY error into NOT_CHECKED, dropping evidence already found."""
    original = agent.apply_guardrails

    def old_rule(model_decision, verdict, errors, ok_searches, seen, guard_flag=None):
        if errors:
            return "NOT_CHECKED", "a search or model call failed; item could not be verified"
        return original(model_decision, verdict, errors, ok_searches, seen)
    monkeypatch.setattr(agent, "apply_guardrails", old_rule)
    report = run_offline([BY_ID["T07"], BY_ID["T08"]])
    assert report["invariants"]["SAFE-4"]["violations"] > 0
    assert {v["case"] for v in report["violations"]} == {"T07", "T08"}


def test_a_guardrail_that_over_blocks_correct_behavior_is_also_caught(monkeypatch):
    monkeypatch.setattr(agent, "apply_guardrails", lambda *a, **k: ("NEEDS_REVIEW", "always escalate"))
    assert run_offline(CASES)["invariants"]["LIVE-1"]["violations"] > 0


def test_every_model_behavior_in_the_matrix_exercises_a_distinct_path():
    finals = run_offline([BY_ID["H01"]])["finals_by_proposal"]
    assert finals["correct_hold"] == {"HOLD": 1} and finals["clear"] == {"NEEDS_REVIEW": 1}
    assert finals["hold_fabricated_recall"] == {"NEEDS_REVIEW": 1} and finals["malformed_fields"] == {"NEEDS_REVIEW": 1}
    assert finals["no_decision"] == {"NEEDS_REVIEW": 1}


def test_the_prompt_injection_case_cannot_change_the_outcome():
    result = run_one(BY_ID["H09"], ScriptedClient(BY_ID["H09"].searches, proposals(BY_ID["H09"])["clear"]))
    assert result["decision"] == "NEEDS_REVIEW" and result["decision"] != "CLEAR"  # a compromised CLEAR is not honored
    hold = run_one(BY_ID["H09"], ScriptedClient(BY_ID["H09"].searches, proposals(BY_ID["H09"])["correct_hold"]))
    assert hold["decision"] == "HOLD"


# ------------------------------------------------------------------ the guardrail fix, directly

SEEN = {"F-1"}
HOLD_OK = {"decision": "HOLD", "confidence": 0.9, "recall_numbers": ["F-1"]}


def test_error_with_a_found_recall_no_longer_hides_it():
    assert apply_guardrails(HOLD_OK, "strong", 1, 1, SEEN) == ("HOLD", None)
    assert apply_guardrails({"decision": "CLEAR", "confidence": 1}, "possible", 1, 1, SEEN)[0] == "NEEDS_REVIEW"
    assert apply_guardrails(None, "strong", 1, 1, SEEN)[0] == "NEEDS_REVIEW"  # e.g. the model call failed after the find


def test_error_with_nothing_found_is_still_not_checked():
    assert apply_guardrails({"decision": "CLEAR", "confidence": 1}, "none", 1, 1, set())[0] == "NOT_CHECKED"
    assert apply_guardrails({"decision": "CLEAR", "confidence": 1}, "none", 0, 0, set())[0] == "NOT_CHECKED"


def test_an_incomplete_check_is_flagged_on_the_result_the_workbook_and_the_email():
    result = run_one(BY_ID["T07"], ScriptedClient(BY_ID["T07"].searches, proposals(BY_ID["T07"])["correct_hold"]))
    assert result["decision"] == "HOLD" and result["incomplete"] is True and result["seen_recalls"] == ["F-27-2026"]
    clean = run_one(BY_ID["H01"], ScriptedClient(BY_ID["H01"].searches, proposals(BY_ID["H01"])["correct_hold"]))
    assert clean["incomplete"] is False

    wb = load_workbook(io.BytesIO(build_workbook(load_bundled_template(), "2026-10-05", [result])))
    row = wb["Recall Check Log"]
    text = next(row.cell(r, 18).value for r in range(16, 20) if row.cell(r, 5).value)
    assert "this check is incomplete" in text

    review = run_one(BY_ID["T08"], ScriptedClient(BY_ID["T08"].searches, proposals(BY_ID["T08"])["review"]))
    _, body = summary_email("2026-10-05", [review], [], [], "s3://b/k")
    assert "check incomplete: a query failed" in body


# ------------------------------------------------------------------ live scoring (scripted stand-ins for a real model)

REPEATS = 2


def policy(kind):
    def build(case: Case):
        p = proposals(case)
        name = {"oracle": {"HOLD": "correct_hold", "CLEAR": "clear", "NEEDS_REVIEW": "review", "NOT_CHECKED": "clear"}[case.expected],
                "lazy": "clear", "reckless": "hold_cites_real_but_unconfirmed"}[kind]
        return ScriptedClient(case.searches, p[name])
    return build


def live(kind):
    runs = run_live(None, "scripted", CASES, REPEATS, client_for=policy(kind))
    return runs, score_live(runs, CASES)


def test_an_oracle_model_scores_perfectly_with_no_unsafe_proposals():
    _, s = live("oracle")
    assert s["accuracy"] == 1.0 and s["metrics"].get("false_clear", 0) == 0 and s["metrics"].get("guardrail_blocked_correct", 0) == 0
    assert s["by_category"]["tool_failure"]["correct"] == s["by_category"]["tool_failure"]["n"]


def test_a_lazy_always_clear_model_is_caught_by_the_guardrails_but_scores_poorly():
    _, s = live("lazy")
    m, unsafe = s["metrics"], sum(c.expected != "CLEAR" for c in CASES) * REPEATS
    assert m.get("false_clear", 0) == 0 and m.get("false_hold", 0) == 0          # nothing unsafe reaches the record
    assert m["raw_unsafe"] == unsafe and m["raw_unsafe_caught"] == unsafe        # but the model itself was unsafe every time
    assert s["accuracy"] < 0.7 and m["over_escalated"] == sum(c.expected == "HOLD" for c in CASES) * REPEATS


def test_a_reckless_always_hold_model_is_caught_too():
    _, s = live("reckless")
    assert s["metrics"].get("false_hold", 0) == 0 and s["metrics"]["raw_unsafe_caught"] == s["metrics"]["raw_unsafe"] > 0


def test_token_and_call_accounting_comes_through():
    runs, s = live("oracle")
    one = next(r for r in runs if r["case"] == "H01")
    assert one["usage"] == {"input_tokens": 200, "output_tokens": 40, "model_calls": 2}
    assert s["avg_model_calls"] > 1 and s["avg_input_tokens"] > 100


def test_a_crashing_run_is_recorded_as_a_crash_not_fatal(monkeypatch):
    from evals import harness

    def boom(case, client, model_id="x", guardrail=None):
        raise RuntimeError("unexpected")
    monkeypatch.setattr(harness, "run_one", boom)
    runs = run_live(None, "m", CASES[:2], 1)
    assert [r["final"] for r in runs] == ["CRASH", "CRASH"] and "RuntimeError" in runs[0]["error"]
    assert score_live(runs, CASES[:2])["metrics"]["crash"] == 2


def test_a_search_miss_is_distinguished_from_a_guardrail_failure():
    """A false CLEAR where the model never surfaced the recall is a search failure, reported separately."""
    case = BY_ID["H01"]
    miss = [{"case": "H01", "repeat": 0, "final": "CLEAR", "raw": "CLEAR", "override": None, "seen": [],
             "usage": {"input_tokens": 1, "output_tokens": 1, "model_calls": 1}, "searches": 1, "error": None, "latency_s": 0.1}]
    s = score_live(miss, [case])
    assert s["metrics"]["false_clear"] == 1 and s["metrics"]["false_clear_search_miss"] == 1


def test_wilson_interval():
    assert wilson(0, 0) == (0.0, 1.0)
    lo, hi = wilson(29, 29)
    assert 0.85 < lo < 0.9 and hi == pytest.approx(1.0)
    lo, hi = wilson(1, 2)
    assert lo < 0.5 < hi and abs((lo + hi) / 2 - 0.5) < 0.05


def test_reports_render():
    _, s = live("oracle")
    text = render_live("some-model", s, REPEATS)
    assert "some-model" in text and "95% CI" in text and "adjudicated by a coordinator" in text
    assert "Total violations: 0" in render_offline(run_offline(CASES[:3]))


# ------------------------------------------------------------------ CLI

def test_cli_offline_writes_reports_and_exits_zero(tmp_path):
    out = subprocess.run([sys.executable, "-m", "evals.run", "offline", "--only", "H01,T07", "--out", str(tmp_path)],
                         capture_output=True, text=True, cwd=os.path.join(os.path.dirname(__file__), ".."))
    assert out.returncode == 0, out.stderr
    assert "Total violations: 0" in (tmp_path / "offline-latest.md").read_text()
    assert json.loads((tmp_path / "offline-latest.json").read_text())["runs"] == 2 * len(proposals(CASES[0]))


def test_cli_exits_nonzero_when_an_invariant_is_violated(tmp_path, monkeypatch):
    monkeypatch.setattr(agent, "apply_guardrails", lambda *a, **k: ("CLEAR", None))
    assert eval_cli.main(["offline", "--out", str(tmp_path)]) == 1


def test_cli_rejects_unknown_case_ids(tmp_path):
    with pytest.raises(SystemExit) as e:
        eval_cli.main(["offline", "--only", "NOPE", "--out", str(tmp_path)])
    assert "NOPE" in str(e.value)


# ------------------------------------------------------------------ search coverage (what the first live eval found)

def test_the_eval_catches_a_weak_search_plan_exactly_where_the_live_run_missed(monkeypatch):
    """Mutation test: a plan that only searches the exact name reproduces the live run's misses."""
    from recall_agent.queryplan import Query
    monkeypatch.setattr(agent, "build_plan", lambda item: [Query("recalling_firm" if item.kind == "supplier" else "product_description", item.name, "exact only")])
    report = run_offline(CASES)
    missed = {v["case"] for v in report["violations"] if v["invariant"] == "COVER-1"}
    assert {"H02", "H04", "R02", "R03", "R06"} <= missed  # the five search misses from the live eval


def test_the_eval_catches_a_missing_plan_entirely(monkeypatch):
    monkeypatch.setattr(agent, "build_plan", lambda item: [])
    report = run_offline(CASES)
    assert report["invariants"]["COVER-1"]["violations"] == report["invariants"]["COVER-1"]["applicable"] > 0


def test_every_positive_case_is_found_by_code_alone():
    for case in CASES:
        if case.expected in ("HOLD", "NEEDS_REVIEW"):
            r = run_one(case, ScriptedClient([], None))
            assert r["matcher"]["verdict"] in ("strong", "possible"), case.id
            if case.expected == "HOLD":
                assert r["matcher"]["verdict"] == "strong" and set(case.cite) <= set(r["seen_recalls"]), case.id


def test_an_empty_handed_model_still_yields_a_safe_final_decision():
    finals = run_offline([BY_ID["H01"], BY_ID["C01"], BY_ID["T01"]])["finals_by_proposal"]["empty_handed"]
    assert finals == {"NEEDS_REVIEW": 2, "NOT_CHECKED": 1}  # no decision is never silently a clear or a hold


# ------------------------------------------------------------------ the T06 fix: faulted cases only count when the fault fired

def test_t06_is_now_deterministic_because_the_code_issues_the_faulted_query():
    from evals.harness import fault_hit
    for proposal in ("clear", "correct_hold", "review", "no_decision"):
        r = run_one(BY_ID["T06"], ScriptedClient(BY_ID["T06"].searches, proposals(BY_ID["T06"])[proposal]))
        assert r["decision"] == "NOT_CHECKED" and fault_hit(BY_ID["T06"], r)
    quiet = run_one(BY_ID["T06"], ScriptedClient([], proposals(BY_ID["T06"])["clear"]))
    assert quiet["decision"] == "NOT_CHECKED"  # even a model that searches nothing hits it, via the plan's UPC queries


def test_every_faulted_case_actually_hits_its_fault_even_with_an_empty_handed_model():
    from evals.harness import fault_hit
    for case in CASES:
        if case.fault:
            assert fault_hit(case, run_one(case, ScriptedClient([], None))), case.id


def test_a_run_whose_fault_never_fired_is_not_scored_and_is_reported():
    row = lambda rep, hit: {"case": "T06", "repeat": rep, "final": "CLEAR", "raw": "CLEAR", "override": None, "seen": [],
                            "usage": {"input_tokens": 1, "output_tokens": 1, "model_calls": 1}, "searches": 1,
                            "error": None, "latency_s": 0.1, "fault_hit": hit}
    s = score_live([row(0, False), row(1, True)], CASES)
    assert s["runs"] == 1 and s["skipped"] == [("T06", 0)]
    assert s["metrics"]["false_clear"] == 1  # only the run where the fault fired was scored (and CLEAR there is wrong)
    assert "Not scored: 1 run(s)" in render_live("m", s, 2)
    assert score_live([row(0, False)], CASES)["runs"] == 0  # all skipped: no division by zero
