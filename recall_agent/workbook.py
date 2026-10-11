"""Fill the coordinator's weekly recall-check workbook.

The coordinator's own file is the template. We load it, fill a copy, and add one sheet.
Nothing is regenerated, so their formatting, dropdowns and other sheets survive.

Targets are found by label and header text, not by cell address, so reordered columns or
an edited layout still work. If something we need is missing we raise WorkbookError and the
handler falls back to the markdown/JSON record and says so in the run notes. We do not
guess a cell.

Rules the writer follows:
  - One row per recall identified. HOLD and NEEDS_REVIEW get a row per matched recall;
    NOT_CHECKED items get a row too, because the form has no other place to say
    "could not verify". CLEAR items have no row; they appear on the audit sheet.
  - "No Applicable Recalls Found" is filled only when there is no HOLD, NEEDS_REVIEW or
    NOT_CHECKED item. Otherwise it is left blank.
  - Attribution is honest: the agent signs as "Recall Agent (automated)" on rows where it
    made the decision (a HOLD whose notice was delivered). Rows that need a human decision
    leave Reviewer blank, and the weekly Signoff (Name & Date) is always left for a person.
  - Values for dropdown columns must come from the template's own lists.
  - SKU, UPC and quantities come from the inventory lines for the lots that matched the recall.
    A quantity is filled only if every matched lot has one (never an under-count); 0 is a real value.
    Customer impact and disposition are human judgment and stay blank.
  - More rows than the table holds: fill it, end with a "+N more" row, and never tick the box.
"""
from __future__ import annotations

import io
from copy import copy
from datetime import date, timedelta
from pathlib import Path

from . import openfda
from .matcher import lot_norm
from .record import SCOPE

AGENT_NAME = "Recall Agent (automated)"
LOG_SHEET = "Recall Check Log"
AUDIT_SHEET = "Agent Audit Trail"
SOURCES_NOTE = ("openFDA Food Enforcement API only (the data behind the FDA Enforcement Report). "
                "USDA FSIS and supplier notifications were not checked by the agent.")
TEMPLATE_PATH = Path(__file__).parent / "templates" / "Weekly_Supplier_Recall_Check_Log.xlsx"

WRITTEN_HEADERS = (
    "Date Checked", "Source / URL", "Source Checked & Filter Used", "Recall Notice ID / Event ID",
    "Product Name & Brand", "Supplier / Manufacturer", "Lot / Batch / Date Codes", "Hazard / Reason for Recall",
    "Recall Classification", "Action Taken", "Applicability Determination", "Evidence / Attachment Link",
    "Reviewer (Name & Date)", "Status", "Affected SKUs", "UPC / Product Code", "Qty Received", "Qty On Hand",
)
AUDIT_FIRST_ROW = 5
AUDIT_HEADERS = ("Item", "Kind", "Decision", "Why (override or note)", "Queries", "Matcher verdict",
                 "Recalls matched", "Model decision", "Model confidence", "Model reasoning", "Hold notice")


class WorkbookError(Exception):
    pass


def load_bundled_template() -> bytes:
    return TEMPLATE_PATH.read_bytes()


# --------------------------------------------------------------------------- layout

def _layout(ws):
    hdr = next((r for r in range(1, 80) if ws.cell(r, 2).value == "Date Checked"), None)
    if hdr is None:
        raise WorkbookError("could not find the 'Date Checked' header row")
    cols = {str(c.value).strip(): c.column for c in ws[hdr] if c.value}
    missing = [h for h in WRITTEN_HEADERS if h not in cols]
    if missing:
        raise WorkbookError(f"template is missing column(s): {missing}")
    notes = next((r for r in range(hdr + 1, ws.max_row + 1)
                  if str(ws.cell(r, 2).value or "").strip() == "Notes:"), None)
    if notes is None:
        raise WorkbookError("could not find the 'Notes:' row that ends the entry table")
    labels = {str(ws.cell(r, 3).value).strip(): r for r in range(1, hdr)
              if isinstance(ws.cell(r, 3).value, str)}
    return hdr, cols, notes, labels


def _dropdowns(ws, cols) -> dict[str, list[str]]:
    by_col = {c: h for h, c in cols.items()}
    out = {}
    for dv in ws.data_validations.dataValidation:
        if dv.type == "list" and dv.formula1:
            vals = [v.strip() for v in dv.formula1.strip('"').split(",")]
            for rng in dv.sqref.ranges:
                if rng.min_col in by_col:
                    out[by_col[rng.min_col]] = vals
    return out


def _clip_validations(ws, last_row: int) -> None:
    """The template's dropdown ranges run past the table into the Notes block; end them at the table."""
    from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
    for dv in ws.data_validations.dataValidation:
        clipped = [CellRange(min_col=r.min_col, min_row=r.min_row, max_col=r.max_col,
                             max_row=max(r.min_row, min(r.max_row, last_row))) for r in dv.sqref.ranges]
        dv.sqref = MultiCellRange(" ".join(r.coord for r in clipped))


# --------------------------------------------------------------------------- row content

def _notice_for(item: str, notices: list[dict]) -> dict | None:
    return next((n for n in notices if n["item"] == item), None)


def _delivered(notice: dict | None) -> bool:
    return bool(notice) and notice["status"] in ("sent", "already_notified")


def _queries(result: dict) -> str:
    calls = result["tool_calls"]
    if not calls:
        return "no query completed"
    parts = []
    for c in calls:
        s = f'{c["field"]}="{c["term"]}": {c["status"]}'
        if c.get("attempts", 1) > 1:
            s += f' after {c["attempts"]} attempts'
        if c.get("error"):
            s += f' ({c["error"]})'
        if c.get("truncated"):
            s += " (results truncated)"
        parts.append(s)
    return "; ".join(parts)


def _source_url(result: dict) -> str:
    calls = result["tool_calls"]
    hit = next((c for c in calls if c["status"] == "matches"), None) or (calls[0] if calls else None)
    return (hit or {}).get("url", "")


def _stock(result: dict, m: dict) -> dict:
    """SKU/UPC/quantities for the inventory lots that matched this recall."""
    wanted = {lot_norm(c) for c in m["lot_overlap"]}
    infos = [i for i in result.get("inventory", []) if i["code"] and lot_norm(i["code"]) in wanted]

    def joined(k):
        return ", ".join(dict.fromkeys(i[k] for i in infos if i[k])) or None

    def total(k):
        vals = [i[k] for i in infos]
        return round(sum(vals), 6) if vals and all(v is not None for v in vals) else None

    return {"Affected SKUs": joined("sku"), "UPC / Product Code": joined("upc"),
            "Qty Received": total("qty_received"), "Qty On Hand": total("qty_on_hand")}


def _determination(result: dict, m: dict | None, notice: dict | None) -> str:
    d = result["decision"]
    md = result["model_decision"] or {}
    reasoning = md.get("reasoning", "")
    incomplete = "\nNOTE: at least one query failed, so this check is incomplete." if result.get("incomplete") else ""
    if d == "HOLD":
        lots = ", ".join(m["lot_overlap"]) if m else ""
        out = [f"APPLICABLE: recalled firm/product matches inventory; lot(s) on hand affected: {lots}."]
        if m and m.get("upc_match"):
            out.append(f"Inventory UPC {m['upc_match']} appears in the recall text.")
        out += [f"Agent confidence {float(md.get('confidence', 0)):.2f}. {reasoning}".strip()]
        if notice:
            to = notice["to"]
            out.append({"sent": f"Hold notice sent to {to}.",
                        "already_notified": f"Hold notice already sent earlier to {to}.",
                        "failed": f"HOLD NOTICE NOT DELIVERED ({notice['error']}). Contact the owner manually.",
                        "no_contact": "HOLD NOTICE NOT SENT: no contact on file. Contact the owner manually."
                        }[notice["status"]])
        return "\n".join(out) + incomplete
    if d == "NEEDS_REVIEW":
        why = result["override_reason"] or reasoning
        extra = f"\nAgent reasoning: {reasoning}" if result["override_reason"] and reasoning else ""
        return f"NEEDS HUMAN REVIEW: {why}{extra}{incomplete}"
    errs = [c["error"] for c in result["tool_calls"] if c.get("error")]
    if result["model_error"]:
        errs.append(result["model_error"])
    return f"NOT CHECKED, could not verify: {result['override_reason']}" + (f" ({'; '.join(errs)})" if errs else "")


def _rows_for(result: dict, notice: dict | None, run_date: str, audit_row: int, classes: list[str]) -> list[dict]:
    d = result["decision"]
    pairs = list(zip(result["matcher"]["matches"], result["evidence"]))
    if d == "HOLD":
        pairs = [(m, r) for m, r in pairs if m["level"] == "strong"]
    elif d == "NOT_CHECKED":
        pairs = []
    decided = d == "HOLD" and _delivered(notice)
    common = {
        "Date Checked": run_date,
        "Source / URL": _source_url(result),
        "Source Checked & Filter Used": (f"openFDA food enforcement; {_queries(result)}; "
                                         f"report date within the last {openfda.DEFAULT_DAYS_BACK} days"),
        "Action Taken": "Hold" if decided else "",
        "Evidence / Attachment Link": ("LINK", audit_row),
        "Reviewer (Name & Date)": f"{AGENT_NAME}, {run_date}" if decided else "",
        "Status": "Open",
    }
    rows = []
    for m, rec in pairs:
        cls = rec.get("classification")
        if cls not in classes:
            cls = "Not Yet Classified" if "Not Yet Classified" in classes else ""
        ids = rec["recall_number"] + (f" / {rec['event_id']}" if rec.get("event_id") else "")
        stock = _stock(result, m) if d == "HOLD" else {}
        rows.append({**common, "Recall Notice ID / Event ID": ids, **stock,
                     "Product Name & Brand": rec.get("product_description", ""),
                     "Supplier / Manufacturer": rec.get("recalling_firm", ""),
                     "Lot / Batch / Date Codes": rec.get("code_info", ""),
                     "Hazard / Reason for Recall": rec.get("reason_for_recall", ""),
                     "Recall Classification": cls,
                     "Applicability Determination": _determination(result, m, notice)})
        if m.get("upc_match") and not rows[-1].get("UPC / Product Code"):
            rows[-1]["UPC / Product Code"] = m["upc_match"]
    if not rows:  # NOT_CHECKED, or a review item with no candidate recall
        is_supplier = result["kind"] == "supplier"
        rows.append({**common,
                     "Recall Notice ID / Event ID": "NOT CHECKED" if d == "NOT_CHECKED" else "None identified",
                     "Supplier / Manufacturer": result["item"] if is_supplier else "",
                     "Product Name & Brand": "" if is_supplier else result["item"],
                     "Applicability Determination": _determination(result, None, notice)})
    return rows


# --------------------------------------------------------------------------- audit sheet

def _audit_sheet(wb, style_src, run_date, results, notices, notes) -> None:
    from openpyxl.styles import Alignment, Font
    ws = wb.create_sheet(AUDIT_SHEET)
    ws["A1"] = f"Agent Audit Trail, {run_date}"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)
    ws["A2"] = SCOPE
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A3"] = "SKUs, UPCs and quantities come from the inventory file supplied for this run."
    ws["A3"].font = Font(name="Calibri", size=10, italic=True)
    ws.merge_cells("A2:K2")
    ws.row_dimensions[2].height = 48
    hdr_style = style_src["header"]
    for i, h in enumerate(AUDIT_HEADERS, start=1):
        c = ws.cell(AUDIT_FIRST_ROW - 1, i, h)
        c._style = copy(hdr_style)
    for n, r in enumerate(results):
        md = r["model_decision"] or {}
        notice = _notice_for(r["item"], notices)
        vals = [r["item"], r["kind"], r["decision"], r["override_reason"] or "", _queries(r),
                r["matcher"]["verdict"], ", ".join(m["recall_number"] for m in r["matcher"]["matches"]),
                md.get("decision", ""), md.get("confidence", ""), md.get("reasoning", ""),
                notice["status"] if notice else ""]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(AUDIT_FIRST_ROW + n, i, v)
            c._style = copy(style_src["body"])
            c.alignment = Alignment(wrap_text=True, vertical="top")
    row = AUDIT_FIRST_ROW + len(results) + 1
    if notes:
        ws.cell(row, 1, "Run notes").font = Font(name="Calibri", size=11, bold=True)
        for n in notes:
            row += 1
            ws.cell(row, 1, n).alignment = Alignment(wrap_text=False)
    for col, w in zip("ABCDEFGHIJK", (26, 11, 14, 38, 52, 14, 22, 14, 12, 52, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"A{AUDIT_FIRST_ROW}"


# --------------------------------------------------------------------------- entry point

def build_workbook(template: bytes, run_date: str, results: list[dict], notices: list[dict] = (),
                   notes: list[str] = ()) -> bytes:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.hyperlink import Hyperlink

    notices = list(notices)
    wb = load_workbook(io.BytesIO(template))
    if LOG_SHEET not in wb.sheetnames:
        raise WorkbookError(f"template has no '{LOG_SHEET}' sheet")
    ws = wb[LOG_SHEET]
    hdr, cols, notes_row, labels = _layout(ws)
    first, last = hdr + 1, notes_row - 1
    capacity = last - first + 1
    lists = _dropdowns(ws, cols)

    # ---- header block
    def put(label, value):
        r = labels.get(label)
        if r is None:
            raise WorkbookError(f"template is missing the '{label}' field")
        ws.cell(r, 4).value = value or None  # top-left cell of the merged input range

    needs_attention = [r for r in results if r["decision"] in ("HOLD", "NEEDS_REVIEW", "NOT_CHECKED")]
    put("Check Date:", run_date)
    put("Checked By:", AGENT_NAME)
    put("Next Review Date:", (date.fromisoformat(run_date) + timedelta(days=7)).isoformat())
    put("No Applicable Recalls Found:", "" if needs_attention else f"Yes, {AGENT_NAME}")
    # "Signoff (Name & Date)" is deliberately never written: that is the human's attestation.
    src_row = next((r for t, r in labels.items() if t.startswith("FDA Enforcement Report")), None)
    if src_row:
        ws.cell(src_row, 4).value = SOURCES_NOTE

    # ---- entry rows
    order = {"HOLD": 0, "NEEDS_REVIEW": 1, "NOT_CHECKED": 2, "CLEAR": 3}
    indexed = sorted(enumerate(results), key=lambda t: (order[t[1]["decision"]], t[0]))
    rows = []
    for idx, r in indexed:
        if r["decision"] != "CLEAR":
            rows += _rows_for(r, _notice_for(r["item"], notices), run_date, AUDIT_FIRST_ROW + idx,
                              lists.get("Recall Classification", []))
    if len(rows) > capacity:
        hidden = len(rows) - (capacity - 1)
        rows = rows[:capacity - 1] + [{
            "Date Checked": run_date, "Recall Notice ID / Event ID": f"+{hidden} MORE",
            "Applicability Determination": f"{hidden} more entries do not fit in this table. See the "
                                           f"'{AUDIT_SHEET}' sheet for every item.", "Status": "Open"}]

    for n, row in enumerate(rows):
        r = first + n
        for header, value in row.items():
            c = ws.cell(r, cols[header])
            if header in lists and value and value not in lists[header]:
                raise WorkbookError(f"'{value}' is not in the template's {header} dropdown {lists[header]}")
            if isinstance(value, tuple):  # internal link to the audit sheet row
                target = f"'{AUDIT_SHEET}'!A{value[1]}"
                c.value = f"Agent Audit Trail, row {value[1]}"
                c.hyperlink = Hyperlink(ref=c.coordinate, location=target, display=c.value)
                c.font = Font(name="Calibri", size=11, color="0563C1", underline="single")
            else:
                c.value = None if value in ("", None) else value  # 0 is a real quantity
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = None  # let the row grow to fit wrapped text

    _clip_validations(ws, last)

    _audit_sheet(wb, {"header": ws.cell(hdr, cols["Date Checked"])._style,
                      "body": ws.cell(first, cols["Date Checked"])._style},
                 run_date, results, notices, list(notes))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
